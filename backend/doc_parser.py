"""Document skills: parse PDF/Word/Excel/txt to text, and generate PDF from text.

文档解析与 PDF 生成，作为 Agent 的「技能层」底层实现。
文本切分优先使用 langchain-text-splitters（LangChain 生态），不可用时回退到简单切分。
"""
import html
import io
import logging
import os
from typing import List

logger = logging.getLogger("agent.doc")

_TEXT_EXT = {".txt", ".md", ".csv", ".log", ".json", ".xml", ".html", ".py", ".java", ".kt"}


def parse_document(content: bytes, filename: str) -> str:
    """根据文件扩展名解析文档为纯文本。"""
    if not content:
        return ""
    ext = os.path.splitext(filename or "")[1].lower()
    if ext == ".pdf":
        return _parse_pdf(content)
    if ext == ".docx":
        return _parse_docx(content)
    if ext in (".xlsx", ".xlsm"):
        return _parse_xlsx(content)
    if ext in _TEXT_EXT:
        return content.decode("utf-8", "replace")
    # 兜底：尝试按文本解码
    return content.decode("utf-8", "replace")


def _parse_pdf(content: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(content))
    parts: List[str] = []
    for i, page in enumerate(reader.pages):
        try:
            text = page.extract_text() or ""
        except Exception as e:
            logger.warning("PDF 第 %d 页提取失败: %s", i + 1, e)
            text = ""
        parts.append(text)
    return "\n\n".join(parts).strip()


def _parse_docx(content: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(content))
    parts: List[str] = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def _parse_xlsx(content: bytes) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    parts: List[str] = []
    for ws in wb.worksheets:
        parts.append(f"[工作表: {ws.title}]")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None and str(c).strip()]
            if cells:
                parts.append("\t".join(cells))
        parts.append("")
    return "\n".join(parts).strip()


def split_chunks(text: str, chunk_size: int = 800, overlap: int = 100) -> List[str]:
    """切分长文本。优先使用 langchain-text-splitters，回退到简单切分。"""
    text = (text or "").strip()
    if not text:
        return []
    try:
        from langchain_text_splitters import RecursiveCharacterTextSplitter
        splitter = RecursiveCharacterTextSplitter(chunk_size=chunk_size, chunk_overlap=overlap)
        chunks = splitter.split_text(text)
        return [c.strip() for c in chunks if c.strip()]
    except Exception as e:
        logger.warning("langchain 切分不可用，回退简单切分: %s", e)
    return _simple_split(text, chunk_size, overlap)


def _simple_split(text: str, chunk_size: int, overlap: int) -> List[str]:
    paragraphs = [p.strip() for p in text.split("\n") if p.strip()]
    chunks: List[str] = []
    buf = ""
    for para in paragraphs:
        if len(buf) + len(para) + 1 <= chunk_size:
            buf = (buf + "\n" + para).strip() if buf else para
        else:
            if buf:
                chunks.append(buf)
            buf = para
            while len(buf) > chunk_size:
                chunks.append(buf[:chunk_size])
                buf = buf[chunk_size - overlap:]
    if buf:
        chunks.append(buf)
    return chunks


def text_to_pdf(text: str) -> bytes:
    """将纯文本生成为 PDF（支持中文，使用 reportlab 内置 CID 字体 STSong-Light）。"""
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.lib.units import cm

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    style = ParagraphStyle(
        name="cn",
        fontName="STSong-Light",
        fontSize=12,
        leading=20,
        wordWrap="CJK",
    )
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        topMargin=2 * cm, bottomMargin=2 * cm,
        leftMargin=2 * cm, rightMargin=2 * cm,
    )
    flowables = []
    for line in (text or "").split("\n"):
        if not line.strip():
            flowables.append(Spacer(1, 8))
        else:
            flowables.append(Paragraph(html.escape(line), style))
    doc.build(flowables)
    return buf.getvalue()